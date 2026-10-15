import argparse
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

parser = argparse.ArgumentParser(
    description="Build the six page editable PawTwin English submission document"
)
parser.add_argument("--output", type=Path, required=True, help="Destination DOCX path")
OUT = parser.parse_args().output.resolve()
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = Inches(0.72)
sec.left_margin = sec.right_margin = Inches(0.78)
sec.header_distance = sec.footer_distance = Inches(0.3)
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(0, 0, 0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.05
for name, size in [("Title", 26), ("Subtitle", 12), ("Heading 1", 16), ("Heading 2", 12)]:
    st = doc.styles[name]
    st.font.name = "Calibri"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_after = Pt(6)
    st.paragraph_format.space_before = Pt(11 if name.startswith("Heading") else 0)
    if name.startswith("Heading"):
        st.paragraph_format.keep_with_next = True
        st.font.bold = True
for name in ["List Bullet", "List Number"]:
    doc.styles[name].font.name = "Calibri"
    doc.styles[name].font.size = Pt(10.5)
    doc.styles[name].paragraph_format.space_after = Pt(4)
doc.core_properties.title = "PawTwin hackathon submission"
doc.core_properties.subject = "English project description judge guide and submission checklist"
doc.core_properties.author = ""
doc.core_properties.last_modified_by = ""
doc.core_properties.keywords = "PawTwin Alexa simulation local prototype"
for node in list(doc.styles.element.iter(qn("w:pBdr"))):
    node.getparent().remove(node)
for node in list(doc.element.iter(qn("w:pBdr"))):
    node.getparent().remove(node)

foot = sec.footer.paragraphs[0]
foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = foot.add_run("PawTwin  |  ")
r.font.size = Pt(8)
fld = OxmlElement("w:fldSimple")
fld.set(qn("w:instr"), "PAGE")
foot._p.append(fld)


def p(text, style=None):
    return doc.add_paragraph(text, style)


def h(text):
    return p(text, "Heading 1")


def sub(text):
    return p(text, "Heading 2")


def bullet(text):
    return p(text, "List Bullet")


def numbered(text):
    global step_number
    step_number += 1
    pp = p(f"{step_number}. {text}")
    pp.paragraph_format.left_indent = Inches(0.25)
    pp.paragraph_format.first_line_indent = Inches(-0.25)
    pp.paragraph_format.space_after = Pt(4)
    return pp


def page(title):
    global step_number
    step_number = 0
    pp = h(title)
    pp.paragraph_format.page_break_before = True


def table(headers, rows, widths):
    tb = doc.add_table(rows=1, cols=len(headers))
    tb.autofit = False
    for i, w in enumerate(widths):
        tb.columns[i].width = Inches(w)
    for c, t in zip(tb.rows[0].cells, headers, strict=True):
        c.text = t
    for row in rows:
        for c, t in zip(tb.add_row().cells, row, strict=True):
            c.text = t
    for ridx, row in enumerate(tb.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if ridx == 0:
            repeat = OxmlElement("w:tblHeader")
            trPr.append(repeat)
        for i, c in enumerate(row.cells):
            c.width = Inches(widths[i])
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = c._tc.get_or_add_tcPr()
            marg = OxmlElement("w:tcMar")
            for side, val in [("top", 75), ("left", 100), ("bottom", 75), ("right", 100)]:
                e = OxmlElement("w:" + side)
                e.set(qn("w:w"), str(val))
                e.set(qn("w:type"), "dxa")
                marg.append(e)
            tcPr.append(marg)
            bord = OxmlElement("w:tcBorders")
            for side in ["top", "left", "bottom", "right"]:
                e = OxmlElement("w:" + side)
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "4")
                e.set(qn("w:color"), "D9D9D9")
                bord.append(e)
            tcPr.append(bord)
            shade = OxmlElement("w:shd")
            shade.set(
                qn("w:fill"), "193B57" if ridx == 0 else ("F1F5F8" if ridx % 2 == 0 else "FFFFFF")
            )
            tcPr.append(shade)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after = Pt(0)
                pp.paragraph_format.line_spacing = 1.0
                for rr in pp.runs:
                    rr.font.size = Pt(9.2)
                    rr.font.color.rgb = RGBColor(255, 255, 255) if ridx == 0 else RGBColor(0, 0, 0)
                    rr.font.bold = ridx == 0
    p("")
    return tb


# Page 1
step_number = 0
p("PawTwin hackathon submission", "Title")
p("English project description and judge guide", "Subtitle")
p("15 October 2026  |  Audited source baseline 6342b96")
p(
    "PawTwin helps a family share one traceable care record for each pet. Confirmed "
    "preferences, manual care tasks and measured growth connect across the household App "
    "and a labelled Alexa text simulator. The working demonstration uses local rules and "
    "SQLite; it has no demonstrated live Alexa, AWS or physical device connection."
)
sub("Inspiration")
p(
    "Pet care is a handoff between people. One person remembers a preference, another "
    "handles playtime, and someone else prepares information for a visit. Separate messages"
    " make it difficult to see which record is current and who actually completed a task."
)
p(
    "Mochi is a fictional household pet used to explore that problem. A family member "
    "records that Mochi likes a feather toy and rests after meals. Another member consults "
    "the same confirmed preference, takes over a care task and leaves an accountable "
    "completion record. This story illustrates a product hypothesis; it is not a real "
    "household trial or a measured customer outcome."
)
sub("What it does")
p(
    "Independent local accounts create or join a household through controlled invitations. "
    "Today, Care, Growth, Twin and Family connect several pet profiles, shared records and "
    "permission roles. A display label such as Dad does not change a member's role."
)
p(
    "Confirmed preferences supply the same per-pet context to the App and simulator. "
    "Answers expose source record IDs. Growth reports calculate changes from timestamped "
    "records and preserve missing values. A revocable caregiver link limits access to one "
    "pet, and a visit card exports saved information without adding a diagnosis or "
    "inventing measurements."
)
sub("The family care story")
numbered("The owner saves a preference as a candidate and explicitly confirms it.")
numbered("A second member asks a new simulator question and checks the cited source.")
numbered(
    "A request for playtime produces a draft. The owner selects a time, prepares the "
    "proposal and explicitly confirms the manual-care arrangement."
)
numbered(
    "The second member claims the task and records its completion. A daily report includes "
    "that manual record and its sources; it does not independently verify a physical "
    "action."
)
p(
    "The product website uses labelled fictional examples. Its illustrated screens do not "
    "read private household records. The actual App and simulator use authenticated local "
    "records."
)

# Page 2
page("How we built it")
p(
    "The prototype uses Python 3.12, FastAPI/Uvicorn, Pydantic, SQLite and a "
    "React/TypeScript/Vite frontend. It targets one application instance. Server-side "
    "sessions, CSRF checks, current membership, roles and pet scope govern access. Version "
    "and expiry checks reject stale care confirmations. A background worker produces "
    "persistent reports in the household's time zone."
)
p(
    "Codex assisted bounded implementation, test execution and independent review, as "
    "recorded in the development evidence. The family assistant is local_rules, and its "
    "memory is owner-confirmed SQLite data. An answer or stored preference grants no "
    "permission and cannot override a hard device limit."
)
sub("Capability matrix")
table(
    ["Capability", "Working demonstration", "Integration scope"],
    [
        (
            "Accounts and records",
            "Local accounts, roles and multiple pets",
            "Single-instance SQLite; no cloud identity",
        ),
        (
            "Confirmed memory",
            "Per-pet preferences with review status and sources",
            "Application data; no AgentCore extraction",
        ),
        (
            "Manual care",
            "Draft, confirmation, claim and recorded completion",
            "No family task actuates a device",
        ),
        (
            "Growth and reports",
            "Timestamped records and deterministic calculations",
            "Demo samples labelled; no inferred intake",
        ),
        (
            "Caregiver link",
            "One-pet read-only access, expiry and revocation",
            "No write, membership or device permission",
        ),
        (
            "Visit information card",
            "Preview, print and offline HTML export",
            "Saved facts; exported copies remain static",
        ),
        (
            "Alexa text simulator",
            "Typed requests and page-scoped conversation",
            "No live Alexa endpoint or microphone input",
        ),
        (
            "Original Mochi workbench",
            "Approval, hard limits and receipt states",
            "Simulated sensors and actuator",
        ),
        (
            "Optional AWS paths",
            "Wiring and offline adapter checks",
            "Live Bedrock, AgentCore and DynamoDB unverified",
        ),
    ],
    [1.45, 2.6, 2.89],
)
sub("Verification and its limits")
p(
    "The 14 October audit records 447 backend tests passed, 1 skipped and 21 live AWS tests"
    " deselected. The full browser suite passed 115 tests across four independent services "
    "with zero retries. Ruff, generated-contract checks, TypeScript and the production "
    "build passed. A 15 October backend rerun returned the same 447 / 1 / 21 result in "
    "52.49 seconds; the browser and static/build figures remain the 14 October evidence."
)
p(
    "The skipped test is test_kill_a_running_server_mid_analysis_then_relaunch. "
    "Process-reopen checks do not replace that unexecuted service-kill case. "
    "Windows/Chromium tests establish the covered local behavior. Remote Linux CI, "
    "production security, model quality, physical hardware and real household outcomes are "
    "not demonstrated."
)
p(
    "Evidence in the source repository: docs/reviews/current-version-2026-10-14.md, its "
    "verify.txt and e2e.txt logs, English-interface evidence dated 14 October, and "
    "caregiver/visit-card evidence dated 12 to 13 October."
)

# Page 3
page("Judge setup and walkthrough")
p(
    "Use Windows/PowerShell with Git, Python 3.12, uv and Node.js/npm. Node 24 is "
    "recommended. Initial dependency and test-browser installation needs network access. "
    "AWS credentials and Docker are not required for the local workflow."
)
sub("Build and launch")
for line in [
    "./scripts/setup.ps1",
    "./scripts/preflight.ps1",
    "./scripts/verify.ps1",
    "./scripts/dev-home.ps1 -Port 8860 -DatabasePath ./data/judge-demo-20261015.db",
]:
    pp = p(line)
    pp.paragraph_format.space_after = Pt(2)
    for rr in pp.runs:
        rr.font.name = "Consolas"
        rr.font.size = Pt(9)
p(
    "Run from the project root and keep the server terminal open. verify.ps1 builds the "
    "frontend. Choose a new disposable database filename and an unused port; do not reset "
    "an existing database. Open http://127.0.0.1:8860/#app/home. The interface defaults to "
    "English, with an English / 中文 switch. User-entered names and notes retain their "
    "original language."
)
p(
    "Other local entries are / for the product website, /#alexa for the simulator and "
    "/#workbench for the separate simulated device workflow. The launcher binds to "
    "127.0.0.1, pins local_rules and local_fake, disables photo recognition and "
    "proxy-header rewriting. Family data uses the selected database path plus .family.db. "
    "Records persist; a server restart requires login again."
)
sub("Test with two independent accounts")
numbered(
    "Use two browser profiles or contexts. Register the owner, create a household and "
    "select Mochi. Generate a carer invitation in Family; the second account joins with it."
    " Two tabs in one profile share cookies and do not demonstrate two users."
)
numbered(
    "Save Likes a feather toy and rests after meals as candidate memory in Twin. Ask What "
    "does it like in the simulator before and after owner confirmation. Only the confirmed "
    "record should be cited. Inspect its source ID and local_rules label."
)
numbered(
    "Enter plan playtime. Choose a concrete due time, prepare the proposal, then confirm "
    "manual care. Draft and preparation create no task. Confirmation creates one manual "
    "task and does not execute a device."
)
numbered(
    "In the carer account, claim the task and record its completion for this fictional "
    "example. Generate the owner's daily report for the household-local date. In a fresh "
    "demonstration, expect one manual completion record with its source."
)
numbered(
    "Delete the confirmed memory and ask a new question. The new answer must stop citing "
    "it. Earlier replies remain historical snapshots. Optionally test a short-lived "
    "caregiver link, revoke it and revisit, then preview/export a visit card."
)
p(
    "Use only fictional staging records and label demo measurements. Passwords, invitation "
    "tokens and caregiver keys should stay out of public footage. Natural-language dates "
    "are not parsed; choose the time in the draft. Changed or expired context requires a "
    "new request."
)
sub("Automated browser check")
p(
    "Set PAWTWIN_E2E_PORT to an unused port such as 8861, then run ./scripts/test-e2e.ps1 "
    "-InstallBrowser. Omit the switch after Chromium is installed, or explicitly select an "
    "existing compatible executable. Four fresh test services isolate registration limits "
    "and keep the judge database separate. The standalone judge-guide.en.md contains the "
    "complete commands, optional weight example and evidence links."
)

# Page 4
page("Development results and next steps")
sub("Challenges")
p(
    "Shared care required careful handling of asynchronous state. Review reproduced delayed"
    " records arriving after an account switch, invitation state surviving logout and "
    "unsaved measurements following a pet switch. The repairs invalidate old responses and "
    "reset protected drafts at the relevant account or pet boundary. A reading timeout now "
    "permits recovery without inventing a task outcome."
)
p(
    "Confirmations also had to refer to the records actually reviewed. Server-held proposal"
    " versions and expiry checks reject stale confirmation. Report regressions distinguish "
    "scheduling from completion, preserve source and author history, and check retry after "
    "failed generation. English interface support preserves user-written notes instead of "
    "treating them as translatable interface labels."
)
sub("Accomplishments")
p(
    "The local prototype connects confirmed memory, an independent second account, a care "
    "draft, explicit confirmation, manual completion and a source-backed report. Caregiver "
    "access and the visit card extend the same records into a controlled handoff. The "
    "audited source includes the complete family App, independent simulator, workbench, "
    "product website and English-first interface."
)
p(
    "Local regressions cover permissions, account and pet isolation, revoked access, stale "
    "proposals and simulated receipt outcomes. The 14 October audit closed the reproducible"
    " review findings and passed the 115-test browser suite. Optional cloud adapters remain"
    " labelled separately from the demonstrated local product."
)
sub("What we learned")
p(
    "A useful pet context needs an identifiable source and a clear review status. A "
    "preference becomes usable after confirmation. An answer does not grant authority. A "
    "care task becomes completed when an authenticated member records the action. Food "
    "offered, dispensing and actual intake need separate meanings; missing intake remains "
    "unknown."
)
p(
    "Negative cases shaped the workflow: switching accounts during a delayed read, changing"
    " a source before confirmation, revoking caregiver access and receiving an unknown "
    "device receipt. A successful demonstration needs those boundaries to remain visible, "
    "especially when a fictional story makes the flow easy to follow."
)
sub("What is next")
p(
    "We want to evaluate the workflow with consenting households before making adoption or "
    "care-outcome claims. Technical next steps include reducing the frontend bundle, "
    "checking accessibility across more devices and browsers, and independently validating "
    "any chosen cloud/model path. Real Alexa and physical device work would need its own "
    "integration and safety evidence."
)
p(
    "The proposed submission route is an Alexa+ simulated experience built with "
    "Codex-assisted development. The shown runtime remains local_rules and SQLite. Final "
    "track selection remains the participant's decision; acceptance is subject to organizer"
    " review. This package does not establish personal eligibility."
)
sub("Source evidence")
bullet(
    "The two-account rehearsal is in frontend/e2e/demo-rehearsal.spec.ts and checks source "
    "IDs, the actual completing user and manual-care provenance."
)
bullet(
    "The latest source review and raw verification logs are in docs/reviews and "
    "docs/evidence/current-version-2026-10-14."
)
bullet(
    "New standalone English copy is in docs/submission/2026-10-15. Older submission "
    "documents and videos remain historical."
)

# Page 5
page("Product Feedback participant review draft")
p(
    "This draft describes recorded project use and developer friction. The contributor must"
    " confirm the tool list, add first-hand onboarding observations and give Yes/No reuse "
    "reasons for every used tool, API or SDK. Personal opinions and real-user feedback are "
    "not invented here."
)
table(
    ["Recorded tools", "Purpose and observed result", "Limit or participant field"],
    [
        (
            "Codex",
            "Bounded implementation, local checks and independent review",
            "No inferred productivity figure or live runtime AI; personal reuse pending",
        ),
        (
            "Python 3.12; FastAPI; Uvicorn",
            "Backend/API, validation paths and same-origin serving",
            "Windows-local evidence; deployment and production proxy use unverified",
        ),
        (
            "Pydantic; pydantic-settings",
            "Models, input validation and generated contracts",
            "Null text and inherited .env required fixes; schema alone cannot enforce workflow",
        ),
        (
            "SQLite; sqlite3; zoneinfo; tzdata",
            "Persistence, reports and household date windows",
            "No distributed proof; UTC/local boundaries required regression coverage",
        ),
        (
            "hashlib.scrypt; secrets",
            "Password hashes and tokens",
            "No complete production security audit; restart requires login",
        ),
        (
            "React; React DOM; TypeScript",
            "Forms, shared cards and typed API use",
            (
                "Delayed scope changes required browser regressions; personal learning "
                "experience pending"
            ),
        ),
        (
            "Vite; React plugin; Node.js; npm",
            "Locked frontend install and production build",
            "Compatible Node required; main chunk exceeds 500 kB",
        ),
        (
            "Playwright; Chromium",
            "Two-account and API-backed browser checks",
            "Fresh shards isolate rate limits; traces can contain session data",
        ),
        (
            "pytest; HTTPX; Ruff",
            "Behavior regressions and Python static checks",
            "1 skipped and 21 live AWS deselected; static success is not cloud proof",
        ),
        (
            "PowerShell; uv; Git",
            "Windows setup, frozen installs and audited history",
            "Remote CI/source access unverified; personal reuse reasons pending",
        ),
        (
            "NumPy; Pillow; SAPI; Media Foundation; FFmpeg",
            "Video composition, offline English narration, encoding/decoding",
            "No live Alexa voice; human playback/listening still needs review",
        ),
        (
            "python-docx; Word; Poppler",
            "Editable document and actual page rendering",
            "LibreOffice renderer unavailable; personal onboarding/reuse pending",
        ),
    ],
    [1.65, 2.56, 2.73],
)
sub("Recorded friction and suggestions")
p(
    "Review repaired delayed account responses, invitation state after logout and cross-pet"
    " drafts. Fixtures now ignore developer .env values, and source export excludes traces "
    "and databases. Four fresh browser-test services isolate registration limits without "
    "weakening production policy. These are local implementation/test issues; the dated "
    "reviews do not report Amazon service failures."
)
sub("Complete before external use")
bullet(
    "Confirm actual use, including final video tools. Give first-hand onboarding detail and"
    " Yes/No reuse reasons individually, or explicitly for every named tool in a group. See"
    " product-feedback.en.md for the full fields."
)
bullet(
    "Keep Bedrock/Boto3, Strands, AgentCore and DynamoDB outside the live-service list "
    "without dated actual-use evidence. Cloud quality, latency and onboarding are "
    "unmeasured."
)

# Page 6
page("Submission checklist and participant fields")
p(
    "The published deadline is 23 October 2026 at 12:00 PDT, equivalent to 24 October 2026 "
    "at 03:00 Asia/Shanghai. Recheck the submission portal before the final submission. The"
    " proposed track is the Alexa+ simulated experience built with Codex-assisted "
    "development; the shown runtime remains local_rules and SQLite."
)
sub("Required external steps")
table(
    ["Submission field", "Current state", "Required action"],
    [
        ("Devpost entry and team", "Pending", "Select track and confirm contributor/team details"),
        (
            "Eligibility and asset rights",
            "Not established",
            "Participant confirms personal eligibility and source/media rights",
        ),
        (
            "GitHub repository",
            "URL not supplied",
            "Share complete runnable source, assets and setup instructions",
        ),
        (
            "Repository permissions",
            "Pending",
            (
                "Public: choose an open-source license. Private: grant required free reviewer "
                "access through judging"
            ),
        ),
        (
            "Public demonstration video",
            "URL not supplied",
            "Upload English video under 3 minutes to YouTube or Vimeo; verify public access",
        ),
        ("English description", "Local copy prepared", "Review and paste the final Devpost copy"),
        (
            "Product Feedback",
            "Participant-review draft",
            "Confirm actual use, onboarding and Yes/No reuse reasons for every used tool",
        ),
        (
            "AI or agentic tool declaration",
            "Codex evidence recorded",
            "Describe development tool separately from the local_rules runtime",
        ),
        (
            "Clean setup and access check",
            "Local evidence available",
            "Test the exact shared repository/package and signed-out video link",
        ),
        (
            "Formal submission",
            "Not performed",
            "Review and submit the Devpost entry before the deadline",
        ),
    ],
    [1.65, 1.55, 3.74],
)
sub("Review the final package")
bullet(
    "Use the same source version for the description, video and setup. The new 15 October "
    "video is a separate production artifact whose final checks belong to its production "
    "record."
)
bullet(
    "Preserve fictional, Demo sample, local_rules, simulator and manual-care labels. SQLite"
    " memory is not AgentCore extraction; a simulated receipt is not hardware operation."
)
bullet(
    "Keep passwords, active tokens, session traces, .env files, databases and backups out "
    "of public artifacts. Missing measurements and intake remain unknown; visit information"
    " is not a diagnosis."
)
p(
    "The simulation route does not require Alexa hardware or public App hosting. Live AWS "
    "is not a mandatory local-simulation prerequisite. An AWS Builder bonus claim needs "
    "actual qualifying use; offline adapters do not establish it. No publish, upload, cloud"
    " deployment or competition submission is completed by this document."
)
sub("Official references")
p("Rules: https://amazonappdev2026.devpost.com/rules")
p("Overview: https://amazonappdev2026.devpost.com/")
p("FAQ: https://amazonappdev2026.devpost.com/details/faqs")
p(
    "Standalone English materials: project-description.en.md, judge-guide.en.md, "
    "product-feedback.en.md and submission-checklist.md in docs/submission/2026-10-15."
)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(str(OUT))
print("paragraphs", len(doc.paragraphs), "tables", len(doc.tables))
