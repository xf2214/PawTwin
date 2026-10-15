"""Build a new editable final document without changing the original delivery."""

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[4]
MATERIALS = ROOT / "docs/submission/2026-10-15"
REFERENCE = ROOT.parent / "PawTwin-submission-2026-10-15.en.docx"
DEFAULT_OUTPUT = ROOT.parent / "PawTwin-submission-2026-10-15.final.en.docx"
FILM = ROOT / "docs/evidence/submission-video-2026-10-15/fullfilm-01"


def build(output: Path, replace_own_output: bool = False) -> None:
    output = output.resolve()
    if output == REFERENCE.resolve():
        raise ValueError("The original DOCX is preserve-only")
    if output.exists() and not replace_own_output:
        raise FileExistsError(f"Refusing to replace an existing output: {output}")
    verification = json.loads((FILM / "verification.json").read_text(encoding="utf-8"))
    for name in [
        "README.md",
        "project-description.en.md",
        "judge-guide.en.md",
        "submission-checklist.md",
        "final-video-production.md",
    ]:
        if not (MATERIALS / name).is_file():
            raise FileNotFoundError(MATERIALS / name)

    # Retain the source document's section geometry, styles, theme and page footer.
    doc = Document(REFERENCE)
    for child in list(doc.element.body):
        if child.tag != qn("w:sectPr"):
            doc.element.body.remove(child)
    doc.core_properties.title = "PawTwin final English hackathon submission"
    doc.core_properties.subject = (
        "Project description judge guide film record and participant fields"
    )
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = "PawTwin local prototype Alexa text simulator final film"
    for node in list(doc.element.iter(qn("w:pBdr"))):
        node.getparent().remove(node)
    for node in list(doc.styles.element.iter(qn("w:pBdr"))):
        node.getparent().remove(node)
    for name in ["Title", "Subtitle", "Heading 1", "Heading 2"]:
        doc.styles[name].font.color.rgb = RGBColor(0, 0, 0)

    def p(text: str, style: str | None = None):
        return doc.add_paragraph(text, style)

    def sub(text: str):
        return p(text, "Heading 2")

    def page(text: str):
        paragraph = p(text, "Heading 1")
        paragraph.paragraph_format.page_break_before = True

    def bullet(text: str):
        return p(text, "List Bullet")

    def step(number: int, text: str):
        paragraph = p(f"{number}. {text}")
        paragraph.paragraph_format.left_indent = Inches(0.25)
        paragraph.paragraph_format.first_line_indent = Inches(-0.25)
        paragraph.paragraph_format.space_after = Pt(4)

    def code(text: str):
        paragraph = p(text)
        paragraph.paragraph_format.space_after = Pt(2)
        for run in paragraph.runs:
            run.font.name = "Consolas"
            run.font.size = Pt(9)

    def table(headers, rows, widths):
        result = doc.add_table(rows=1, cols=len(headers))
        result.autofit = False
        for index, width in enumerate(widths):
            result.columns[index].width = Inches(width)
        for cell, text in zip(result.rows[0].cells, headers, strict=True):
            cell.text = text
        for values in rows:
            for cell, text in zip(result.add_row().cells, values, strict=True):
                cell.text = text
        for row_index, row in enumerate(result.rows):
            properties = row._tr.get_or_add_trPr()
            properties.append(OxmlElement("w:cantSplit"))
            if row_index == 0:
                properties.append(OxmlElement("w:tblHeader"))
            for index, cell in enumerate(row.cells):
                cell.width = Inches(widths[index])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                properties = cell._tc.get_or_add_tcPr()
                margins = OxmlElement("w:tcMar")
                for side, value in [("top", 75), ("left", 100), ("bottom", 75), ("right", 100)]:
                    item = OxmlElement("w:" + side)
                    item.set(qn("w:w"), str(value))
                    item.set(qn("w:type"), "dxa")
                    margins.append(item)
                properties.append(margins)
                borders = OxmlElement("w:tcBorders")
                for side in ["top", "left", "bottom", "right"]:
                    item = OxmlElement("w:" + side)
                    item.set(qn("w:val"), "single")
                    item.set(qn("w:sz"), "4")
                    item.set(qn("w:color"), "D9D9D9")
                    borders.append(item)
                properties.append(borders)
                shade = OxmlElement("w:shd")
                shade.set(
                    qn("w:fill"),
                    "193B57" if row_index == 0 else ("F1F5F8" if row_index % 2 == 0 else "FFFFFF"),
                )
                properties.append(shade)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    for run in paragraph.runs:
                        run.font.size = Pt(9.2)
                        run.font.color.rgb = (
                            RGBColor(255, 255, 255) if row_index == 0 else RGBColor(0, 0, 0)
                        )
                        run.font.bold = row_index == 0
        p("")
        return result

    # Page 1: project description and actual family story.
    p("PawTwin final English hackathon submission", "Title")
    p("Project description judge guide and participant fields", "Subtitle")
    p("15 October 2026  |  Final local source identity recorded in release-source.json")
    p(
        "PawTwin helps a family share one traceable care record for each pet. Confirmed "
        "preferences, manual care tasks and measured growth connect across the household App "
        "and a labelled Alexa text simulator. This package describes the local prototype and "
        "the 142.379-second final English film. Publication and formal submission remain pending."
    )
    sub("Inspiration")
    p(
        "Pet care is often a handoff. One person remembers a preference, another handles "
        "playtime, and someone else prepares information for a visit. Separate messages make "
        "it difficult to see which record is current and who recorded a task's completion. "
        "Mochi, Emma and Alex are fictional examples used to explore that product hypothesis. "
        "They are not a household trial or evidence of measured customer outcomes."
    )
    sub("What the working App does")
    p(
        "Independent local accounts create or join a household through controlled invitations. "
        "Today, Care, Growth, Twin and Family connect several pet profiles, permissions and "
        "shared records. Display names such as Dad do not change a member's permission role. "
        "Confirmed preferences supply per-pet context to the App and simulator, with source IDs "
        "available for inspection. Reports calculate from saved records "
        "and preserve unknown values."
    )
    sub("The family care sequence in the final film")
    step(
        1,
        "Emma saves Mochi's feather-toy and after-meal rest preference as candidate memory, "
        "reviews it and explicitly confirms it.",
    )
    step(
        2,
        "Alex asks what Mochi likes in a fresh Alexa text simulator conversation. The reply "
        "uses the confirmed preference and exposes its source.",
    )
    step(
        3,
        "Emma requests a playtime draft, chooses a time and prepares a proposal. Only her "
        "explicit confirmation creates a shared manual-care task.",
    )
    step(
        4,
        "Alex claims the task through his own account and records completion. Emma's daily "
        "report includes the saved care record. This does not independently verify physical care.",
    )
    sub("Scripted Pet Concierge preview")
    p(
        "The film then shows fictional Miso and Luna. Miso's food question is paired with a "
        "preset profile and preference; Luna's illustrative activity records accompany observation "
        "guidance. The product page and workbench show these fixed scenarios without runtime "
        "model inference, actual tracking, diagnosis, private family-data reads or task creation."
    )

    # Page 2: implementation, capability comparison and dated evidence.
    page("Implementation and verified scope")
    p(
        "The single-instance prototype uses Python 3.12, FastAPI/Uvicorn, Pydantic, SQLite and "
        "React/TypeScript/Vite. Server-side sessions, CSRF checks, current membership, roles and "
        "pet scope govern access. Server-held proposal versions and expiry reject stale care "
        "confirmations. A worker generates persistent reports in the household's time zone."
    )
    p(
        "Codex assisted bounded implementation, local test execution and independent review. "
        "The demonstrated family assistant uses local_rules. Its confirmed memory is application "
        "data in SQLite. An answer or preference grants no permission "
        "and cannot override a hard limit."
    )
    table(
        ["Capability", "Working local scope", "Boundary"],
        [
            (
                "Accounts and records",
                "Households, invitations, roles and multiple pets",
                "Single-instance storage; no cloud identity",
            ),
            (
                "Confirmed memory",
                "Reviewed per-pet preferences and source IDs",
                "No demonstrated AgentCore extraction",
            ),
            (
                "Care handoff",
                "Draft, explicit confirmation, claim and manual completion",
                "Recorded care; no device actuation",
            ),
            (
                "Growth and reports",
                "Timestamped measures and deterministic calculations",
                "Unknown intake remains unknown",
            ),
            (
                "Caregiver and visit card",
                "Expiring read-only access and saved-information export",
                "No write/approval authority or diagnosis",
            ),
            (
                "Alexa experience",
                "Typed-request browser simulator",
                "No live endpoint or microphone input",
            ),
            (
                "Pet Concierge",
                "Fictional preset examples and illustrative traces",
                "No model, tracking or diagnosis",
            ),
            (
                "Original workbench",
                "Owner approval, hard limits and simulated receipts",
                "Mochi sensors and actuator are simulated",
            ),
            (
                "Optional AWS paths",
                "Adapter preparation and offline tests",
                "No demonstrated live Bedrock, AgentCore or DynamoDB",
            ),
        ],
        [1.45, 2.70, 2.79],
    )
    sub("Historical audit and final repair evidence")
    p(
        "The historical 14 October audit at 6342b96 passed 447 backend tests, with 1 skipped "
        "and 21 live AWS tests deselected, plus 115 browser tests with zero retries. Ruff, "
        "contracts, TypeScript and production build passed then. That commit predates the "
        "15 October Pet Concierge additions and does not identify the final working tree."
    )
    p(
        "A later 15 October component audit recorded 447 backend passed, 1 skipped, 21 "
        "deselected and 122 browser passed before final repairs. Final whole-tree and package "
        "checks belong in docs/evidence/submission-repairs-2026-10-15/verification.json. Final "
        "source identity belongs in docs/submission/2026-10-15/release-source.json "
        "and the external source manifest, prepared by the finalization step. "
        "Earlier counts do not close those gates."
    )
    p(
        "The skipped service-kill/relaunch test is not a pass. Windows/Chromium checks do not "
        "establish remote Linux CI, live cloud, Alexa hardware or physical care effectiveness."
    )

    # Page 3: setup, workflow and additional application checks.
    page("Judge setup and walkthrough")
    p(
        "Use Windows/PowerShell with Git, Python 3.12, uv and a compatible Node.js/npm runtime; "
        "Node 24 is recommended. First dependency/browser installation needs network access. "
        "The local workflow does not require Docker or AWS credentials."
    )
    sub("Build and launch from the project root")
    for line in [
        "./scripts/setup.ps1",
        "./scripts/preflight.ps1",
        "./scripts/verify.ps1",
        "./scripts/dev-home.ps1 -Port 8860 -DatabasePath ./data/judge-demo-20261015.db",
    ]:
        code(line)
    p(
        "Keep the server terminal open; verify.ps1 builds the frontend. Choose a new disposable "
        "database filename and unused port. Open http://127.0.0.1:8860/#app/home. The interface "
        "defaults to English with an English / 中文 switch; saved user text keeps its language. "
        "The launcher binds to 127.0.0.1, uses local_rules/local_fake, disables photo recognition "
        "and proxy-header rewriting. Family records use the selected path plus .family.db."
    )
    sub("Reproduce the working family handoff")
    step(
        1,
        "Use two independent browser profiles or contexts; ordinary tabs share login cookies. "
        "Register an owner, create a household and select Mochi. Invite a second account as carer.",
    )
    step(
        2,
        "Save Likes a feather toy and rests after meals as candidate memory in Twin. Ask What "
        "does it like before and after owner confirmation. Only the confirmed preference should "
        "be cited. Inspect its source ID and local_rules label in a fresh simulator session.",
    )
    step(
        3,
        "Enter plan playtime, choose a concrete due time, prepare the proposal and explicitly "
        "confirm manual care. Drafting or preparing alone creates no care task.",
    )
    step(
        4,
        "In the second account, claim the task and record completion for the fictional example. "
        "Generate a daily report for the household-local date; in a fresh demonstration expect "
        "one manual completion and its source. A saved completion is not a physical-care sensor.",
    )
    sub("Additional runnable checks outside the final film")
    bullet(
        "Delete confirmed memory and ask again; new answers must stop citing it. Earlier replies "
        "remain snapshots. Test a short-lived caregiver link, revoke it and revisit. Preview or "
        "export a visit card; demo measurements require explicit inclusion and keep their labels."
    )
    bullet(
        "Growth example: Demo sample weights of 1.20 kg on 25 September and 1.35 kg on 2 October "
        "2026 at the same noon time give 150 g, 12.5% and seven elapsed days in the weekly report. "
        "Missing intake stays not recorded. The earlier weight is the comparison baseline."
    )
    bullet(
        "Open / for the product site, /#alexa for the text simulator, /#workbench for the device "
        "demo or /#workbench/concierge for preset Miso/Luna scenarios. Existing workbench demo "
        "login remains required. Keep all credentials and invitation/caregiver tokens private."
    )
    p(
        "For browser checks use an unused PAWTWIN_E2E_PORT "
        "and ./scripts/test-e2e.ps1 -InstallBrowser. "
        "Four fresh services isolate registration limits. See judge-guide.en.md for complete "
        "commands, timing rules, source links and limitations."
    )

    # Page 4: final film identity and meaningful development narrative.
    page("Final film and development results")
    sub("Active English film")
    p(
        "PawTwin-hackathon-submission-2026-10-15.en.mp4 is the active film, accompanied by the "
        "same-stem .srt and .poster.png. It lasts "
        f"{verification['actual_duration_seconds']:.3f} seconds at 1920 by 1080 and 30 fps, "
        "with H.264 video, AAC stereo audio and 51 caption cues."
    )
    p("MP4 SHA256  " + verification["sha256"])
    p(
        "The film follows confirmed memory, Alex's fresh simulator conversation, Emma's explicit "
        "task confirmation, Alex's recorded completion and Emma's daily report, then the scripted "
        "Miso/Luna preview and an honest local boundary. Caregiver access, visit cards, growth "
        "calculations and memory deletion are additional app features, not claimed film scenes."
    )
    sub("Technical review and remaining human review")
    p(
        "The fullfilm-01 verification records a complete stream decode, "
        "48 encoded frames inspected and independent review of 14 critical frames. "
        "The 51 cues match narration with no overlap. "
        "Encoded audio measures -16.01 LUFS and -1.82 dBTP with no clipped samples or truncation. "
        "These are technical checks. Complete human playback/listening remains pending, including "
        "narration clarity, music balance and caption timing. No upload or submission occurred."
    )
    p(
        "The final production uses continuous Ava narration "
        "and original locally synthesized music. Earlier SAPI/no-music "
        "and 150-second v2 tool descriptions apply to historical productions. "
        "Use final-video-production.md and the fullfilm-01 delivery, audio "
        "and verification records "
        "for this film. Media generation does not establish live Alexa or in-product inference."
    )
    sub("Challenges and accomplishments")
    p(
        "Review reproduced delayed account responses, invitation state surviving logout and "
        "measurement drafts following a pet switch. Repairs invalidate stale responses and reset "
        "protected drafts. Server-held versions and expiry keep confirmations tied to reviewed "
        "context. Reports preserve authors and sources and distinguish scheduling from completion."
    )
    p(
        "The local prototype connects confirmed preferences, independent accounts, explicit care "
        "confirmation, authenticated manual completion and a source-backed report. Caregiver "
        "access and the visit card provide additional controlled handoffs. Negative cases, "
        "including revoked access and unknown simulated receipts, shaped the implementation."
    )
    sub("What we learned and what comes next")
    p(
        "Pet context needs identifiable sources and review status. An answer grants no authority; "
        "a recorded completion remains a human account of care. "
        "Offered food, dispensing and intake have separate meanings. "
        "We want consenting household trials before adoption/outcome claims, "
        "plus bundle-size and accessibility work. Any cloud/model, real Alexa or hardware path "
        "requires separate validation."
    )
    p(
        "The proposed route is an Alexa+ simulated experience "
        "built with Codex-assisted development. Track selection, eligibility "
        "and organizer acceptance remain participant/external decisions."
    )

    # Page 5: explicitly incomplete personal feedback without stale final-media claims.
    page("Product Feedback participant completion")
    p(
        "Recorded development use supports this tool summary. The participant must confirm actual "
        "use, add first-hand onboarding experience and answer Yes or No with a reason for every "
        "used tool, API or SDK. No personal opinion, setup time or real-user feedback is invented."
    )
    table(
        ["Recorded tool group", "Observed purpose", "Participant field or limit"],
        [
            (
                "Codex",
                "Bounded implementation, local checks and independent review",
                "Confirm personal onboarding and reuse; no inferred productivity gain",
            ),
            (
                "Python 3.12; FastAPI; Uvicorn",
                "Backend logic, validation and same-origin serving",
                "Windows-local evidence; production/distributed behavior unverified",
            ),
            (
                "Pydantic; pydantic-settings",
                "Input models and generated contracts",
                "Null text and inherited .env required repairs; schema is not workflow authority",
            ),
            (
                "SQLite; sqlite3; zoneinfo; tzdata",
                "Local persistence and household date windows",
                "No distributed-cloud proof; confirm each tool's reuse reason",
            ),
            (
                "hashlib.scrypt; secrets",
                "Password hashing and tokens",
                "No complete production security audit; restart requires login",
            ),
            (
                "React; React DOM; TypeScript",
                "UI forms, typed API use and shared cards",
                "Scope changes required browser regressions; accessibility coverage incomplete",
            ),
            (
                "Vite; React plugin; Node; npm",
                "Locked installs and production build",
                "Compatible Node needed; main chunk exceeds 500 kB",
            ),
            (
                "Playwright; Chromium; pytest; HTTPX; Ruff",
                "API-backed browser, backend and static checks",
                "Skipped/live tests are not passes; traces can contain session data",
            ),
            (
                "PowerShell; uv; Git",
                "Windows setup, frozen installs and source history",
                "Remote CI/reviewer access unverified; personal reuse pending",
            ),
            (
                "Remotion; edge-tts; NumPy; media codecs",
                "Final film composition, continuous Ava narration and original music",
                "Separate production evidence; not live Alexa or runtime model use",
            ),
            (
                "python-docx; Word; Poppler; pdfplumber",
                "Editable document and actual rendered-page QA",
                "Canonical LibreOffice unavailable previously; see new document-final-qa.json",
            ),
        ],
        [1.85, 2.35, 2.74],
    )
    sub("Recorded friction and suggestions")
    p(
        "Account-switch responses, invitation cleanup and cross-pet drafts required regressions. "
        "Fixtures ignore developer .env values; source exports exclude traces and database files. "
        "Fresh test services preserve production registration limits. These are local engineering "
        "issues, not reports of Amazon service failures. Confirm final media tools against their "
        "recorded production versions rather than copying historical feedback unchanged."
    )
    sub("Complete before external use")
    bullet(
        "Confirm every actual-use entry and provide onboarding detail plus Yes/No reuse reasons "
        "for each named tool. product-feedback.en.md contains detailed participant fields "
        "and distinguishes historical media tooling from the current final-film record."
    )
    bullet(
        "Bedrock/Boto3, Strands, AgentCore and DynamoDB remain adapter/offline paths without "
        "dated live-use evidence. Cloud onboarding, latency and model quality are unmeasured; "
        "do not claim an AWS Builder bonus from imports or offline tests."
    )

    # Page 6: finalization outputs and remaining external gates.
    page("Submission checklist and source identity")
    p(
        "The recorded deadline is 23 October 2026 at 12:00 PDT, equivalent to 24 October 2026 at "
        "03:00 Asia/Shanghai. Recheck the official event portal immediately before submission. "
        "This local package does not establish eligibility or organizer acceptance."
    )
    sub("Final local package records")
    p(
        "The finalization step prepares PawTwin-source-2026-10-15.final.zip, "
        "PawTwin-submission-materials-2026-10-15.final.zip and their external JSON manifests. "
        "The new editable document is PawTwin-submission-2026-10-15.final.en.docx. Confirm "
        "actual existence, hashes, archive CRC, links and current film/source identity in the "
        "finalization records. No future commit is invented in this document."
    )
    p(
        "Use release-source.json for the final local source reference and the external source "
        "manifest for the actual freeze commit. "
        "Use submission-repairs-2026-10-15/verification.json "
        "for the repaired whole-tree result, and document-final-qa.json for this DOCX's page QA. "
        "Historical baseline 6342b96 and the original document QA "
        "do not identify these new outputs."
    )
    table(
        ["External field", "Current state", "Required participant action"],
        [
            (
                "Entry track and team",
                "Pending",
                "Select track; verify contributor and team details",
            ),
            (
                "Eligibility and rights",
                "Not established",
                "Confirm residence/age/employment exclusions and source/media/music rights",
            ),
            (
                "Source repository URL",
                "Not supplied",
                "Provide complete runnable source and setup with required reviewer access",
            ),
            (
                "License or private access",
                "Pending",
                "If public choose a license; if private grant prescribed "
                "free access through judging",
            ),
            (
                "Complete human film review",
                "Pending",
                "Watch/listen to all 142.379 seconds; review narration, music and captions",
            ),
            (
                "Public video URL",
                "Not supplied",
                "Upload to YouTube/Vimeo and verify signed-out access",
            ),
            (
                "English description",
                "Local text prepared",
                "Review the final text and paste it into Devpost",
            ),
            (
                "Product Feedback",
                "Participant fields pending",
                "Confirm actual tools, onboarding and individual reuse reasons",
            ),
            (
                "AI tool declaration",
                "Codex evidence recorded",
                "Describe development assistance separately from local_rules runtime",
            ),
            (
                "Final setup and submission",
                "Not performed",
                "Test exact shared source/access, review and submit before the deadline",
            ),
        ],
        [1.70, 1.55, 3.69],
    )
    p(
        "Preserve fictional, Demo sample, local_rules, simulator, manual-care and preset labels. "
        "Keep credentials, live tokens, traces, .env files and databases private. Missing values "
        "remain unknown and visit information is not a diagnosis. The simulation route does not "
        "require Alexa hardware or public App hosting. Cloud, publication and formal submission "
        "are separate actions and are not completed by this local preparation."
    )
    sub("Official references and current copy")
    p(
        "Rules  https://amazonappdev2026.devpost.com/rules\n"
        "Overview  https://amazonappdev2026.devpost.com/\n"
        "FAQ  https://amazonappdev2026.devpost.com/details/faqs"
    )
    p(
        "Current English text and evidence entry points are in docs/submission/2026-10-15: "
        "README.md, project-description.en.md, judge-guide.en.md, submission-checklist.md and "
        "final-video-production.md. Earlier 150/170-second films "
        "and previous documents remain historical."
    )

    output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output)
    print(
        json.dumps(
            {
                "output": str(output),
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables),
                "planned_pages": 6,
            }
        )
    )


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument(
        "--replace-own-output",
        action="store_true",
        help="Allow QA iteration on this new output; original is always protected",
    )
    args = parser.parse_args()
    build(args.output, args.replace_own_output)
